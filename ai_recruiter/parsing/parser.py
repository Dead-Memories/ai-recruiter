"""Извлечение текста из резюме и вакансий.

Поддерживаемые форматы: PDF, DOCX, ODT, TXT. Единая точка входа —
`extract_text(path)`, которая диспетчеризует по расширению файла и
возвращает нормализованный текст.

Запуск для проверки на папке cvs:
    python -m ai_recruiter.parsing.parser cvs
"""

from __future__ import annotations

import argparse
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".odt", ".txt"}

# Пространства имён OpenDocument (для text:p / text:h)
_ODT_NS = {
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


class ParsingError(Exception):
    """Ошибка извлечения текста из файла."""


def extract_text(path: str | Path) -> str:
    """Извлекает и нормализует текст из файла по его расширению."""
    path = Path(path)
    if not path.exists():
        raise ParsingError(f"Файл не найден: {path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(path)
    elif ext == ".docx":
        text = _extract_docx(path)
    elif ext == ".odt":
        text = _extract_odt(path)
    elif ext == ".txt":
        text = _extract_txt(path)
    else:
        raise ParsingError(
            f"Неподдерживаемый формат: {ext!r} (поддерживаются: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))})"
        )

    return normalize_text(text)


def normalize_text(text: str) -> str:
    """Приводит текст к единообразному виду, сохраняя абзацы."""
    text = text.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x0c", "\n")

    lines = [line.strip() for line in text.split("\n")]
    out: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                out.append("")
            blank = True
        else:
            out.append(line)
            blank = False
    return "\n".join(out).strip()


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Не удалось открыть PDF: {path} ({exc})") from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            raise ParsingError(
                f"Ошибка извлечения текста из страницы: {path} ({exc})"
            ) from exc
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Не удалось открыть DOCX: {path} ({exc})") from exc

    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _extract_odt(path: Path) -> str:
    try:
        with zipfile.ZipFile(str(path)) as archive:
            if "content.xml" not in archive.namelist():
                raise ParsingError(f"В ODT-файле нет content.xml: {path}")
            root = ET.fromstring(archive.read("content.xml"))
    except zipfile.BadZipFile as exc:
        raise ParsingError(f"Файл не является ODT-архивом: {path}") from exc

    lines: list[str] = []
    for elem in root.iter():
        tag = elem.tag.split("}")[-1]
        if tag in ("p", "h"):
            lines.append("".join(elem.itertext()))
    return "\n".join(lines)


def _extract_txt(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1251"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ParsingError(f"Не удалось определить кодировку TXT-файла: {path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Извлечение текста из резюме/вакансий")
    parser.add_argument("paths", nargs="+", help="Файлы или папки для обработки")
    parser.add_argument("--max-chars", type=int, default=800, help="Лимит вывода на файл")
    args = parser.parse_args()

    targets: list[Path] = []
    for p in args.paths:
        p = Path(p)
        if p.is_dir():
            targets.extend(sorted(f for f in p.iterdir() if f.suffix.lower() in SUPPORTED_EXTENSIONS))
        else:
            targets.append(p)

    for path in targets:
        try:
            text = extract_text(path)
        except ParsingError as exc:
            print(f"[SKIP] {path.name}: {exc}")
            continue
        preview = text[: args.max_chars]
        print(f"\n===== {path.name} ({len(text)} chars) =====")
        print(preview)


if __name__ == "__main__":
    main()
