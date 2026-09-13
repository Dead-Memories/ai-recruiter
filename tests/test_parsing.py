"""Unit-тесты модуля парсинга (PDF/DOCX/ODT/TXT)."""

from __future__ import annotations

import zipfile

import pytest

from ai_recruiter.parsing import SUPPORTED_EXTENSIONS, ParsingError, extract_text, normalize_text


def _make_pdf(path) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    c.drawString(72, 720, "Hello QA engineer")
    c.showPage()
    c.save()


def _make_docx(path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Python backend developer")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "PostgreSQL"
    table.rows[0].cells[1].text = "Docker"
    doc.save(str(path))


def _make_odt(path) -> None:
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        "<office:body><office:text>"
        "<text:h>Resume</text:h>"
        "<text:p>Data Analyst</text:p>"
        "<text:p>SQL and Python</text:p>"
        "</office:text></office:body>"
        "</office:document-content>"
    )
    with zipfile.ZipFile(str(path), "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)


class TestNormalizeText:
    def test_removes_non_breaking_spaces(self):
        assert normalize_text("a\u00a0b") == "a b"

    def test_collapses_blank_lines(self):
        assert normalize_text("a\n\n\n\nb") == "a\n\nb"

    def test_strips_whitespace(self):
        assert normalize_text("  a  \n  b  ") == "a\nb"


class TestExtractText:
    def test_txt(self, tmp_path):
        f = tmp_path / "resume.txt"
        f.write_text("line one\nline two\n", encoding="utf-8")
        assert extract_text(f) == "line one\nline two"

    def test_txt_cp1251_fallback(self, tmp_path):
        f = tmp_path / "resume.txt"
        f.write_text("Резюме\n", encoding="cp1251")
        assert extract_text(f) == "Резюме"

    def test_docx(self, tmp_path):
        f = tmp_path / "resume.docx"
        _make_docx(f)
        text = extract_text(f)
        assert "Python backend developer" in text
        assert "PostgreSQL" in text

    def test_pdf(self, tmp_path):
        f = tmp_path / "resume.pdf"
        _make_pdf(f)
        assert "Hello QA engineer" in extract_text(f)

    def test_odt(self, tmp_path):
        f = tmp_path / "resume.odt"
        _make_odt(f)
        text = extract_text(f)
        assert "Resume" in text
        assert "Data Analyst" in text

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(ParsingError):
            extract_text(tmp_path / "nope.pdf")

    def test_unsupported_extension_raises(self, tmp_path):
        f = tmp_path / "resume.csv"
        f.write_text("a,b", encoding="utf-8")
        with pytest.raises(ParsingError):
            extract_text(f)


def test_supported_extensions():
    assert {".pdf", ".docx", ".odt", ".txt"} <= SUPPORTED_EXTENSIONS
